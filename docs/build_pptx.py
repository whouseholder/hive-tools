#!/usr/bin/env python3
"""
Build the hive-tools overview slide deck as a native PowerPoint (.pptx).

This is an *authoring-time* helper only -- it is NOT one of the shipped tools and
is never imported by them. It renders the same content as docs/slides.html into a
16:9 deck with native PowerPoint titles, bullets and tables, embedding the Mermaid
diagrams as pre-rendered PNGs from docs/assets/.

Usage (from a machine with the dev dependency installed):

    python3 -m pip install python-pptx        # authoring only
    python3 docs/build_pptx.py                 # writes docs/hive-tools-overview.pptx

The diagram PNGs in docs/assets/ were exported from docs/slides.html (Mermaid).
Re-export them if the diagrams change, then re-run this script.
"""

import os
import re
from collections import namedtuple

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# --------------------------------------------------------------------------- #
# Paths
# --------------------------------------------------------------------------- #
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
OUT = os.path.join(HERE, "hive-tools-overview.pptx")

DG = {
    "problem":      os.path.join(ASSETS, "problem.png"),
    "which_tool":   os.path.join(ASSETS, "which_tool.png"),
    "architecture": os.path.join(ASSETS, "architecture.png"),
    "detection":    os.path.join(ASSETS, "detection.png"),
    "safety":       os.path.join(ASSETS, "safety.png"),
    "pipeline":     os.path.join(ASSETS, "pipeline.png"),
    "workflow":     os.path.join(ASSETS, "workflow.png"),
}

# --------------------------------------------------------------------------- #
# Palette / type (mirrors the accent colours used in docs/slides.html)
# --------------------------------------------------------------------------- #
ACCENT  = RGBColor(0x1D, 0x6F, 0xB8)   # blue   -- titles, links, code
ACCENT2 = RGBColor(0x2E, 0x7D, 0x32)   # green  -- "safe / no mutation"
WARN    = RGBColor(0xB2, 0x3B, 0x3B)   # red    -- "mutating / danger"
INK     = RGBColor(0x1F, 0x29, 0x33)   # body text
MUTED   = RGBColor(0x52, 0x60, 0x6D)   # secondary text
HDRBG   = RGBColor(0xEE, 0xF3, 0xF8)   # table header fill
ROWBG   = RGBColor(0xF7, 0xF9, 0xFB)   # zebra row fill
RULE    = RGBColor(0xD6, 0xDE, 0xE6)   # thin rules
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
CHIPBG  = RGBColor(0xE6, 0xF2, 0xFB)

BODY = "Calibri"
MONO = "Consolas"

COLORS = {"red": WARN, "green": ACCENT2, "muted": MUTED, "accent": ACCENT,
          "ink": INK, "white": WHITE, None: INK}
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}

EMU_IN = 914400
Seg = namedtuple("Seg", "text bold mono")

# --------------------------------------------------------------------------- #
# Presentation (16:9)
# --------------------------------------------------------------------------- #
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
MARGIN = 0.55
CW = SW - 2 * MARGIN            # usable content width
TOP = 1.42                      # content start (below the title rule)


def _rgb(c):
    return c if isinstance(c, RGBColor) else COLORS.get(c, INK)


def _tokenize(text):
    """Split a string into runs on **bold** and `mono` markup."""
    segs, pos = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text):
        if m.start() > pos:
            segs.append(Seg(text[pos:m.start()], False, False))
        if m.group(1) is not None:
            segs.append(Seg(m.group(1), True, False))
        else:
            segs.append(Seg(m.group(2), False, True))
        pos = m.end()
    if pos < len(text):
        segs.append(Seg(text[pos:], False, False))
    return segs or [Seg(text, False, False)]


def _emit(p, text, size, color, bold=False, italic=False):
    """Add markup-aware runs to an existing paragraph."""
    for s in _tokenize(text):
        r = p.add_run()
        r.text = s.text
        f = r.font
        f.size = Pt(size)
        f.name = MONO if s.mono else BODY
        f.bold = bool(s.bold or bold)
        f.italic = italic
        f.color.rgb = ACCENT if s.mono else _rgb(color)


# --------------------------------------------------------------------------- #
# Slide furniture
# --------------------------------------------------------------------------- #
def new_slide():
    return prs.slides.add_slide(BLANK)


def add_rect(slide, left, top, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                 Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_title(slide, title):
    tb = slide.shapes.add_textbox(Inches(MARGIN), Inches(0.34),
                                  Inches(CW), Inches(0.72))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _emit(p, title, 26, ACCENT, bold=True)
    add_rect(slide, MARGIN, 1.12, CW, 0.032, ACCENT)


def add_footer(slide, page, total):
    tb = slide.shapes.add_textbox(Inches(MARGIN), Inches(7.08),
                                  Inches(7), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    _emit(p, "hive-tools \u2022 overview", 9, MUTED)
    tb2 = slide.shapes.add_textbox(Inches(SW - MARGIN - 1.4), Inches(7.08),
                                   Inches(1.4), Inches(0.3))
    p2 = tb2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    _emit(p2, "%d / %d" % (page, total), 9, MUTED)


class Box:
    """Thin wrapper around a text frame for adding markup paragraphs."""

    def __init__(self, slide, left, top, w, h, anchor="top"):
        self.tf = slide.shapes.add_textbox(Inches(left), Inches(top),
                                           Inches(w), Inches(h)).text_frame
        self.tf.word_wrap = True
        self.tf.vertical_anchor = ANCHOR[anchor]
        self._used = False

    def add(self, text, size=15, color=INK, bold=False, italic=False,
            bullet=False, align="left", space_before=0, space_after=7,
            line_spacing=1.12):
        p = self.tf.paragraphs[0] if not self._used else self.tf.add_paragraph()
        self._used = True
        p.alignment = ALIGN[align]
        p.space_after = Pt(space_after)
        p.space_before = Pt(space_before)
        p.line_spacing = line_spacing
        if bullet:
            r = p.add_run()
            r.text = "\u2022   "
            r.font.size = Pt(size)
            r.font.name = BODY
            r.font.color.rgb = ACCENT
        _emit(p, text, size, color, bold=bold, italic=italic)
        return p


def add_image_contain(slide, key, left, top, w, h):
    """Place an image scaled to fit (contain) within a box, centred."""
    with Image.open(DG[key]) as im:
        iw, ih = im.size
    ar = iw / ih
    box_ar = w / h
    if box_ar > ar:                 # box wider -> height bound
        ph = h
        pw = h * ar
    else:                           # width bound
        pw = w
        ph = w / ar
    plx = left + (w - pw) / 2.0
    ply = top + (h - ph) / 2.0
    slide.shapes.add_picture(DG[key], Emu(int(plx * EMU_IN)), Emu(int(ply * EMU_IN)),
                             width=Emu(int(pw * EMU_IN)), height=Emu(int(ph * EMU_IN)))


# --------------------------------------------------------------------------- #
# Tables
# --------------------------------------------------------------------------- #
def P(text, size=12, color=None, small=False, align=None, bold=False):
    return {"text": text, "size": (10.5 if small else size),
            "color": color, "align": align, "bold": bold}


def C(content, span=1, align="left", valign="top", fill=None, color=None,
      size=12, small=False):
    """Normalise a cell spec into a dict of paragraphs + layout options."""
    if isinstance(content, str):
        paras = [P(content, size=size, color=color, small=small, align=align)]
    elif isinstance(content, dict):
        paras = [content]
    else:
        paras = []
        for x in content:
            if isinstance(x, str):
                paras.append(P(x, size=size, color=color, small=small, align=align))
            else:
                paras.append(x)
    return {"paras": paras, "span": span, "align": align, "valign": valign, "fill": fill}


def _set_grid_style(tbl):
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    sid = tblPr.makeelement(qn("a:tableStyleId"), {})
    sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # "No Style, Table Grid"
    tblPr.append(sid)


def _fill_cell(cell, color):
    cell.fill.solid()
    cell.fill.fore_color.rgb = color


def _write_cell(cell, spec, header=False):
    cell.vertical_anchor = ANCHOR[spec.get("valign", "top")]
    cell.margin_left = Inches(0.09)
    cell.margin_right = Inches(0.09)
    cell.margin_top = Inches(0.05)
    cell.margin_bottom = Inches(0.05)
    tf = cell.text_frame
    tf.word_wrap = True
    first = True
    for pr in spec["paras"]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ALIGN[pr.get("align") or spec.get("align") or "left"]
        p.line_spacing = 1.06
        p.space_after = Pt(2)
        color = ACCENT if header else pr.get("color")
        _emit(p, pr["text"], pr.get("size", 12), color, bold=(header or pr.get("bold")))


def add_table(slide, left, top, col_w, header, rows,
              header_size=12, body_size=12, zebra=True):
    ncols = len(col_w)
    nrows = len(rows) + 1
    total_w = sum(col_w)
    gf = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top),
                                Inches(total_w), Inches(0.34 * nrows))
    tbl = gf.table
    _set_grid_style(tbl)
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Inches(w)

    # header
    for c, htext in enumerate(header):
        cell = tbl.cell(0, c)
        _fill_cell(cell, HDRBG)
        _write_cell(cell, C(htext, size=header_size), header=True)

    # body
    for r, row in enumerate(rows, start=1):
        ci = 0
        for spec in row:
            if isinstance(spec, dict) and "paras" in spec:
                cell_spec = spec
            else:
                cell_spec = C(spec, size=body_size)
            span = cell_spec.get("span", 1)
            origin = tbl.cell(r, ci)
            if span > 1:
                origin.merge(tbl.cell(r, ci + span - 1))
            if cell_spec.get("fill"):
                _fill_cell(origin, cell_spec["fill"])
            elif zebra and r % 2 == 0:
                _fill_cell(origin, ROWBG)
            else:
                _fill_cell(origin, WHITE)
            _write_cell(origin, cell_spec)
            ci += span
    return gf


def add_chip(slide, left, top, w, text):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left),
                                 Inches(top), Inches(w), Inches(0.44))
    shp.fill.solid()
    shp.fill.fore_color.rgb = CHIPBG
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _emit(p, text, 13, ACCENT, bold=True)


# =========================================================================== #
# Slides
# =========================================================================== #
def slide_title():
    s = new_slide()
    add_rect(s, 0, 0, SW, 0.16, ACCENT)               # top accent band
    Box(s, 1.5, 1.55, SW - 3.0, 0.4).add(
        "CDP 7.1.9 (Hive 3)  \u2022  on-prem  \u2022  Isilon  \u2022  MySQL metastore",
        size=13, color=MUTED, align="center")
    Box(s, 1.5, 2.05, SW - 3.0, 1.1).add("hive-tools", size=54, color=ACCENT,
                                          bold=True, align="center")
    Box(s, 1.5, 3.25, SW - 3.0, 0.6).add(
        "Keeping the **Hive Metastore** and its **MySQL backend** healthy.",
        size=20, color=INK, align="center")
    add_chip(s, 6.667 - 3.35, 4.15, 2.7, "Orphan Cleanup")
    add_chip(s, 6.667 + 0.35, 4.15, 3.0, "Metadata Pressure Monitor")
    Box(s, 2.4, 5.05, SW - 4.8, 1.2).add(
        "Two standalone Python tools (3.6+, standard library only). This deck: why "
        "they exist, what they do, and their limits \u2014 the repo README is the full guide.",
        size=13, color=MUTED, align="center")
    return s


def slide_why():
    s = new_slide()
    add_title(s, "Why these tools exist")
    b = Box(s, MARGIN, TOP, CW, 1.9)
    b.add("The metastore's MySQL DB is a shared, single-writer bottleneck. "
          "Two things degrade it:", size=15, color=MUTED, space_after=9)
    b.add("**Row bloat** \u2014 data deleted straight off storage leaves **orphaned** "
          "table/partition rows behind in MySQL. Millions of dead rows slow every "
          "plan, backup, and upgrade.", size=15, bullet=True)
    b.add("**Call pressure** \u2014 badly-shaped workloads flood the HMS with API calls; "
          "each is a MySQL round-trip, so metadata gets slow **for everyone**.",
          size=15, bullet=True)
    add_image_contain(s, "problem", MARGIN, 3.55, CW, 3.15)
    return s


def slide_two_tools():
    s = new_slide()
    add_title(s, "The two tools at a glance")
    header = ["Tool", "What it does", "Changes the cluster?"]
    rows = [
        [C("**Orphan Cleanup**  `hive_orphan_cleanup.py`"),
         C("Finds & removes Hive tables/partitions whose **storage is gone**, using "
           "standard Hive DML \u2014 shrinking metastore row counts on MySQL."),
         C([P("Only with `--execute`", color="red"),
            P("read-only by default", small=True, color="muted")])],
        [C("**Pressure Monitor**  `metadata_pressure_monitor.py`"),
         C("Reads HMS logs to find **what/who** is overloading the metastore, "
           "correlates to Hive operations & likely queries/apps, and prescribes fixes."),
         C([P("**Never**", color="green"),
            P("strictly read-only", small=True, color="muted")])],
    ]
    add_table(s, MARGIN, TOP, [3.2, 6.4, 2.63], header, rows)
    Box(s, MARGIN, 5.9, CW, 0.5).add(
        "Both are dependency-free: copy one script to an edge node and run it.",
        size=12, color=MUTED)
    return s


def slide_which_tool():
    s = new_slide()
    add_title(s, "Which tool, when?")
    add_image_contain(s, "which_tool", MARGIN, TOP, 5.7, 5.2)
    b = Box(s, 6.9, TOP + 0.1, 5.85, 5.0)
    b.add("Not the right tool when\u2026", size=17, color=INK, bold=True, space_after=10)
    for t in [
        "You want to tune **MySQL** itself or change Cloudera Manager configs.",
        "You want to find **unused** tables \u2014 orphan = storage genuinely gone, not idle.",
        "You expect it to **apply** fixes \u2014 the monitor only **recommends**.",
        "You want to reclaim **disk** \u2014 cleanup removes metadata, not data.",
    ]:
        b.add(t, size=15, bullet=True, space_after=10)
    return s


def slide_requirements():
    s = new_slide()
    add_title(s, "Requirements & scope")
    header = ["Need", "Orphan Cleanup", "Pressure Monitor"]
    rows = [
        [C("Python"),
         C("3.6+ , **standard library only** (RHEL 8 friendly)", span=2, align="center")],
        [C("Cluster connection"),
         C("`beeline` + `hdfs` on an edge node"),
         C("**None** \u2014 reads log files", color="green")],
        [C("Auth"),
         C("Kerberos ticket (or keytab / principal)"),
         C("**None**", color="green")],
        [C("Platform"),
         C("Built & verified for CDP 7.1.9 / Hive 3 / Isilon / MySQL", span=2, align="center")],
        [C("Install"),
         C("**Nothing** \u2014 no pip, no virtualenv", span=2, align="center", color="green")],
    ]
    add_table(s, MARGIN, TOP, [3.0, 4.615, 4.615], header, rows)
    Box(s, MARGIN, 5.55, CW, 1.0).add(
        "**In scope:** metastore metadata health (row bloat + call pressure).   "
        "**Out of scope:** deleting user data, MySQL tuning, applying configs.",
        size=13, color=MUTED)
    return s


def slide_architecture():
    s = new_slide()
    add_title(s, "Architecture \u2014 systems these tools touch")
    add_image_contain(s, "architecture", MARGIN, TOP - 0.05, CW, 4.55)
    Box(s, MARGIN, 6.05, CW, 1.0).add(
        "**Legend:**  thin arrow = read-only   |   **thick arrow = potentially "
        "mutating** (Orphan Cleanup, only with `--execute`, only via Hive DML)   |   "
        "dotted = optional.  No tool ever issues raw `DELETE` against MySQL.",
        size=12, color=MUTED)
    return s


def slide_orphan_modes():
    s = new_slide()
    add_title(s, "Orphan Cleanup \u2014 what it does")
    Box(s, MARGIN, TOP, CW, 0.5).add(
        "Three modes; the recommended flow is **report \u2192 review \u2192 apply**.",
        size=15, color=MUTED)
    header = ["Mode", "What it does", "Mutates?"]
    rows = [
        [C("`report`"), C("Enumerate HMS, check storage, write CSV + JSON + summary."),
         C("**No**", color="green")],
        [C("`clean`"), C("Detect **and** clean. Dry-run by default."),
         C("With `--execute`")],
        [C("`apply`"), C("Clean exactly the rows in a prior (trimmed) report."),
         C("With `--execute`")],
    ]
    add_table(s, MARGIN, TOP + 0.6, [2.2, 7.4, 2.63], header, rows)
    Box(s, MARGIN, 4.6, CW, 1.4).add(
        "Cleanup is standard Hive DML (`DROP TABLE` / `ALTER TABLE \u2026 DROP PARTITION`), "
        "which cascades away the backing rows in `TBLS` / `PARTITIONS` / `SDS` / \u2026 "
        "\u2014 the MySQL win.", size=13, color=MUTED)
    return s


def slide_detection():
    s = new_slide()
    add_title(s, "Orphan Cleanup \u2014 how detection works")
    add_image_contain(s, "detection", MARGIN, TOP, 5.7, 5.2)
    b = Box(s, 6.9, TOP + 0.1, 5.85, 5.0)
    b.add("Positive proof of absence", size=17, color=INK, bold=True, space_after=10)
    for t in [
        "An object is flagged **only** when its parent lists successfully **and** it "
        "is genuinely missing.",
        "Any **indeterminate** result (outage, NameNode blip) is **skipped** \u2014 a "
        "storage hiccup cannot manufacture orphans.",
        "Listings are batched **one per parent dir** and cached, keeping load off Isilon.",
    ]:
        b.add(t, size=15, bullet=True, space_after=11)
    return s


def slide_safety():
    s = new_slide()
    add_title(s, "Orphan Cleanup \u2014 safety model")
    add_image_contain(s, "safety", MARGIN, TOP, 5.5, 5.35)
    header = ["Gate", "Default", "Loosen with"]
    rows = [
        [C("Dry-run"), C("on"), C("`--execute`")],
        [C("External-only"), C("managed skipped"), C("`--allow-managed`")],
        [C("ACID protection"), C("always"), C("none", color="red")],
        [C("Positive proof"), C("always"), C("discouraged", small=True, color="muted")],
        [C("Bulk guard"), C("on"), C("`--allow-bulk`")],
        [C("Typed confirm"), C("type `yes`"), C("`--yes`")],
    ]
    add_table(s, 6.75, TOP, [2.1, 2.0, 1.9], header, rows, body_size=11.5)
    Box(s, 6.75, 5.35, 6.0, 1.3).add(
        "Loosening one gate never weakens the others: `--yes` skips only the keystroke; "
        "ACID is never droppable.", size=11.5, color=MUTED)
    return s


def slide_monitor_answers():
    s = new_slide()
    add_title(s, "Pressure Monitor \u2014 what it answers")
    b = Box(s, MARGIN, TOP, CW, 4.4)
    for t in [
        "Which HMS **methods** dominate call volume and total time?",
        "Which **Hive operations** (DROP / ALTER / ADD PARTITION, CREATE / DROP TABLE, "
        "ANALYZE) drive them?",
        "Which **users**, **client IPs**, **tables**, and **HMS hosts** contribute most?",
        "How does load **trend over time** \u2014 where are the storms?",
        "Which patterns look like **loops / repetitive calls** \u2014 and how to fix them?",
    ]:
        b.add(t, size=16, bullet=True, space_after=12)
    Box(s, MARGIN, 5.6, CW, 1.0).add(
        "For every finding it prints a concrete remediation: batching, partition "
        "pruning, breaking loops, delegation-token store, stats tuning, \u2026",
        size=13, color=MUTED)
    return s


def slide_pipeline():
    s = new_slide()
    add_title(s, "Pressure Monitor \u2014 how a run works")
    add_image_contain(s, "pipeline", MARGIN, TOP, 5.5, 4.6)
    Box(s, MARGIN, 6.05, 5.5, 0.7).add(
        "Cheap \u2192 expensive, so most lines are discarded before any regex.",
        size=11.5, color=MUTED, align="center")
    header = ["Mode", "Use", "Mutates?"]
    rows = [
        [C("`report`"), C("Analyze a window; fan-out & incremental options."),
         C("**No**", color="green")],
        [C("`monitor`"), C("Scheduled tail with de-duplicated alerts."),
         C("**No**", color="green")],
        [C("`agg`"), C("Per-node worker for SSH fan-out."),
         C("**No**", color="green")],
    ]
    add_table(s, 6.75, TOP, [1.7, 2.95, 1.35], header, rows, body_size=11.5)
    b = Box(s, 6.75, 3.45, 6.0, 3.0)
    for t in [
        "**Fast:** a `--last 1h` scan of a multi-GB log runs in well under a second.",
        "**Incremental** (`--state-dir`): re-runs parse only new bytes.",
        "**Fleet** (`--hosts`): each HMS node parses its own logs.",
    ]:
        b.add(t, size=12.5, bullet=True, space_after=9, color=MUTED)
    return s


def slide_confidence():
    s = new_slide()
    add_title(s, "Pressure Monitor \u2014 attributing the load")
    Box(s, MARGIN, TOP, CW, 1.0).add(
        "Beyond \u201cwhich method is hot\u201d, the report names **who / what to go fix**: "
        "top contributors, and up to a **Top-5 likely source per operation** with a "
        "confidence score.", size=15, color=MUTED)
    header = ["Evidence available", "Confidence", "Shown as"]
    rows = [
        [C("user + table from HMS audit only"),
         C("ranked, capped **< 100%**"),
         C("\u201cNN%\u201d (guess)")],
        [C("+ client op matched on table **and** time"),
         C("higher %"),
         C("\u201cNN%\u201d + queryId / appId")],
        [C("single unambiguous driver"),
         C("**100%**", color="green"),
         C("**CONFIRMED**", color="green")],
    ]
    add_table(s, MARGIN, TOP + 1.15, [5.0, 3.2, 4.03], header, rows)
    Box(s, MARGIN, 5.55, CW, 1.1).add(
        "Honest by design: HMS audit logs share no id with HS2, so attribution is a "
        "ranked **guess** unless HS2 / YARN / Spark logs confirm it. Treat the Top-5 "
        "as leads.", size=13, color=MUTED)
    return s


def slide_implementation():
    s = new_slide()
    add_title(s, "Implementation highlights")
    left = Box(s, MARGIN, TOP, 5.85, 5.0)
    left.add("Built for production", size=17, color=ACCENT, bold=True, space_after=10)
    for t in [
        "Python **3.6+, stdlib only** \u2014 no install.",
        "Verified on **RHEL 8 / CDP 7.1.9**.",
        "Read-only by default; every drop is audited.",
        "Third-party safety **audit** + automated test harness.",
    ]:
        left.add(t, size=15, bullet=True, space_after=11)
    right = Box(s, 6.9, TOP, 5.85, 5.0)
    right.add("Built for repeated runs", size=17, color=ACCENT, bold=True, space_after=10)
    for t in [
        "Fast filtering + cached parsing (**~16\u00d7** faster timestamps).",
        "**Incremental** re-runs (process only new log bytes).",
        "**Fleet fan-out** over SSH \u2014 no central log collection.",
        "Batched, cached `hdfs -ls` to spare Isilon.",
    ]:
        right.add(t, size=15, bullet=True, space_after=11)
    return s


def slide_limitations():
    s = new_slide()
    add_title(s, "Limitations & caveats")
    b = Box(s, MARGIN, TOP, CW, 5.2)
    for t in [
        "**Scope:** tuned for CDP 7.1.9 / Hive 3 on-prem (RHEL 8, Isilon, MySQL). "
        "Other versions untested.",
        "**Orphan Cleanup needs a healthy filesystem view** \u2014 run when Isilon / HDFS "
        "+ NameNode are healthy; the bulk guard is a backstop, not a substitute.",
        "**\u201cOrphaned\u201d \u2260 \u201cunused\u201d** \u2014 tables whose data still exists are never flagged.",
        "**Correlation is best-effort** \u2014 a ranked guess with a confidence score unless "
        "client logs confirm it.",
        "**Logs must be readable & standard** \u2014 parsing targets Cloudera PERFLOG + AUDIT "
        "lines; live tail reads uncompressed logs.",
        "**Change control** \u2014 back up the metastore (MySQL dump) before large cleanups.",
    ]:
        b.add(t, size=15, bullet=True, space_after=12)
    return s


def slide_workflow():
    s = new_slide()
    add_title(s, "A safe end-to-end workflow")
    add_image_contain(s, "workflow", MARGIN, 2.35, CW, 1.6)
    Box(s, 1.4, 4.5, SW - 2.8, 1.6).add(
        "Diagnose with the monitor \u2192 remove dead metadata with cleanup \u2192 confirm the "
        "metastore got faster. Execution stays a human-reviewed step.",
        size=16, color=INK, align="center")
    return s


def slide_learn_more():
    s = new_slide()
    add_title(s, "Where to learn more")
    Box(s, MARGIN, TOP, CW, 0.6).add(
        "The repository **README** files are the comprehensive guide.",
        size=16, color=INK, align="center")
    header = ["Doc", "Contents"]
    rows = [
        [C("`README.md`"), C("Overview, which-tool, systems, safety model, caveats.")],
        [C("`orphan-cleanup/README.md`"), C("Detection, safety gates, options, scheduling.")],
        [C("`pressure-monitor/README.md`"), C("Pipeline, fan-out, correlation, findings catalog.")],
        [C("`AUDIT.md`"), C("Third-party safety & compatibility audit.")],
    ]
    add_table(s, (SW - 11.0) / 2.0, TOP + 0.8, [4.0, 7.0], header, rows)
    Box(s, MARGIN, 5.55, CW, 1.1).add(
        "Start read-only:   `report --paths /var/log/hive --last 24h`    /    "
        "`report --jdbc-url \u2026 --output-dir ./out`",
        size=13, color=MUTED, align="center")
    return s


# --------------------------------------------------------------------------- #
# Assemble
# --------------------------------------------------------------------------- #
def main():
    builders = [
        slide_title, slide_why, slide_two_tools, slide_which_tool,
        slide_requirements, slide_architecture, slide_orphan_modes,
        slide_detection, slide_safety, slide_monitor_answers, slide_pipeline,
        slide_confidence, slide_implementation, slide_limitations,
        slide_workflow, slide_learn_more,
    ]
    for fn in builders:
        fn()

    total = len(prs.slides._sldIdLst)
    for i, slide in enumerate(prs.slides, start=1):
        if i == 1:
            continue                     # no footer on the title slide
        add_footer(slide, i, total)

    prs.save(OUT)
    print("wrote %s  (%d slides)" % (OUT, total))


if __name__ == "__main__":
    main()
